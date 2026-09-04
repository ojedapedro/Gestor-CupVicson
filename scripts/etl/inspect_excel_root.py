import openpyxl, os

base = r"C:/Users/Agencia de viajes/OneDrive/Escritorio/CUP vicson"
files = [
    "023 - VICSON VALENCIA - RECOP 2017 PRECIERRE MAYO.xlsx",
    "CFC+CJC+RG+PXF 2017 vicson valencia.xlsx",
    "CONTROL DE TRASPASOS  VICSON VALENCIA CORTE AL 19 DE MAYO.xls",
    "CUP MAYO 2017 VICSON VALENCIA.xls",
    "GUIAS VICSON 1ERA SEMANA.xls",
    "GUIAS VICSON 2DA SEMANA.xls",
    "INVENTARIO 19 PPRE-CIERRE.xlsx",
]

result = []
for fname in files:
    path = os.path.join(base, fname)
    result.append(f"\n{'='*80}")
    result.append(f"ARCHIVO: {fname}")
    result.append(f"{'='*80}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            result.append(f"\n  --- Hoja: '{sname}' (filas={ws.max_row}, cols={ws.max_column}) ---")
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(10, ws.max_row), values_only=True), 1):
                if any(c is not None for c in row):
                    vals = [str(c)[:80] if c is not None else "" for c in row]
                    result.append(f"    Fila {i}: {vals}")
            if ws.max_row > 10:
                result.append(f"    ... ({ws.max_row - 10} filas más)")
    except Exception as e:
        result.append(f"  ERROR: {e}")

print("\n".join(result))
